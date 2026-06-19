import os
import requests
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from telegram import Update
from telegram.ext import Application, MessageHandler, CommandHandler, filters, ContextTypes
import json
import re
import time

TELEGRAM_TOKEN = os.environ.get("TELEGRAM_TOKEN")
GROQ_KEY = os.environ.get("GROQ_KEY")

# История чатов для каждого пользователя
chat_histories = {}

SYSTEM_PROMPT = """Ты BEK AI — умный, дружелюбный помощник в Telegram. Создан Y.Samir в мае 2026 года.
Ты не Llama, не GPT, не другой AI. Ты только BEK AI.
Правила:
- ВСЕГДА отвечай на том же языке на котором пишет пользователь
- Если пишут по-английски — отвечай по-английски
- Если пишут по-русски — отвечай по-русски
- Если пишут по-арабски — отвечай по-арабски
- Если пишут по-китайски — отвечай по-китайски
- Если пишут по-испански — отвечай по-испански
- Если пишут по-французски — отвечай по-французски
- Если пишут по-немецки — отвечай по-немецки
- Если пишут по-японски — отвечай по-японски
- Никогда не мешай языки в одном ответе
- Понимай короткие и неформальные сообщения
- Отвечай кратко и по делу
- Будь дружелюбным и естественным"""

def create_pptx(slides_data):
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    for slide in slides_data:
        layout = prs.slide_layouts[1]
        s = prs.slides.add_slide(layout)
        background = s.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(10, 10, 30)
        title = s.shapes.title
        body = s.placeholders[1]
        title.text = slide.get("title", "")
        title.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 212, 255)
        title.text_frame.paragraphs[0].font.size = Pt(36)
        title.text_frame.paragraphs[0].font.bold = True
        body.text = slide.get("content", "")
        body.text_frame.paragraphs[0].font.color.rgb = RGBColor(200, 200, 255)
        body.text_frame.paragraphs[0].font.size = Pt(20)
    path = "/tmp/presentation.pptx"
    prs.save(path)
    return path

async def start(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    user_id = update.message.from_user.id
    chat_histories[user_id] = []  # сброс истории
    await update.message.reply_text(
        "👋 Привет! Я BEK AI — умный помощник созданный Y.Samir.\n\n"
        "Напиши /help чтобы узнать что я умею!"
    )

async def help_cmd(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "🤖 BEK AI — возможности:\n\n"
        "💬 Просто напиши мне — отвечу на любой вопрос\n"
        "🖼 фото — пришлю случайное фото\n"
        "📊 презентация [тема] — создам презентацию\n\n"
        "🛠 Поддержка: @Samir_Yl"
    )

async def handle(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    user_text = update.message.text
    user_id = update.message.from_user.id

    # Инициализация истории
    if user_id not in chat_histories:
        chat_histories[user_id] = []

    if user_text.lower().startswith("презентация"):
        topic = user_text[11:].strip() or "тема не указана"
        await update.message.reply_text("⏳ Создаю презентацию...")
        prompt = f"""Создай презентацию на тему: {topic}
Ответь ТОЛЬКО в JSON формате без лишнего текста:
{{"slides": [{{"title": "...", "content": "..."}}, ...]}}
Сделай 10 слайдов."""
        url = "https://api.groq.com/openai/v1/chat/completions"
        headers = {"Authorization": "Bearer " + GROQ_KEY, "Content-Type": "application/json"}
        body = {"model": "llama-3.3-70b-versatile", "messages": [{"role": "user", "content": prompt}]}
        r = requests.post(url, headers=headers, json=body)
        data = r.json()
        try:
            text = data["choices"][0]["message"]["content"]
            match = re.search(r'\{.*\}', text, re.DOTALL)
            if match:
                slides_data = json.loads(match.group())["slides"]
                path = create_pptx(slides_data)
                await update.message.reply_document(document=open(path, "rb"), filename="presentation.pptx")
            else:
                await update.message.reply_text("⚠️ Не смог создать презентацию. Обратитесь в поддержку: @Samir_Yl")
        except Exception:
            await update.message.reply_text("⚠️ Что-то пошло не так. Обратитесь в поддержку: @Samir_Yl")

    elif user_text.lower().startswith("фото"):
        await update.message.reply_text("⏳ Ищу фото...")
        try:
            seed = int(time.time())
            image_url = f"https://picsum.photos/seed/{seed}/512/512"
            img = requests.get(image_url, timeout=30)
            await update.message.reply_photo(photo=img.content)
        except Exception:
            await update.message.reply_text("⚠️ Что-то пошло не так. Обратитесь в поддержку: @Samir_Yl")

    else:
        # Добавляем сообщение пользователя в историю
        chat_histories[user_id].append({"role": "user", "content": user_text})
        
        # Оставляем только последние 10 сообщений
        if len(chat_histories[user_id]) > 10:
            chat_histories[user_id] = chat_histories[user_id][-10:]

        url = "https://api.groq.com/openai/v1/chat/completions"
        headers = {"Authorization": "Bearer " + GROQ_KEY, "Content-Type": "application/json"}
        body = {
            "model": "llama-3.3-70b-versatile",
            "messages": [{"role": "system", "content": SYSTEM_PROMPT}] + chat_histories[user_id]
        }
        r = requests.post(url, headers=headers, json=body)
        data = r.json()
        try:
            answer = data["choices"][0]["message"]["content"]
            # Добавляем ответ бота в историю
            chat_histories[user_id].append({"role": "assistant", "content": answer})
        except Exception:
            answer = "⚠️ Что-то пошло не так. Обратитесь в поддержку: @Samir_Yl"
        await update.message.reply_text(answer)

if __name__ == "__main__":
    app = Application.builder().token(TELEGRAM_TOKEN).build()
    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("help", help_cmd))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle))
    app.run_polling(drop_pending_updates=True)
